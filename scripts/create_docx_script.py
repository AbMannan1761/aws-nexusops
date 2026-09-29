"""
Generates a beautifully formatted Microsoft Word (.docx) document
for the AWS NexusOps 3-Minute Demo Video Script.
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX_PATH = os.path.join(BASE_DIR, "docs", "AWS_NexusOps_Demo_Video_Script.docx")

def build_docx():
    doc = Document()

    # Document Title
    title = doc.add_heading("AWS NexusOps — 3-Minute Demo Video Script", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("AWS Communication Developer Services (CDS) Agentic AI Partner Hackathon 2026")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(11)
    sub.runs[0].font.italic = True
    sub.runs[0].font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Meta Information Box
    meta_p = doc.add_paragraph()
    meta_p.add_run("Target Duration: ").bold = True
    meta_p.add_run("Exactly 2 Minutes 50 Seconds (Hackathon limit: max 3 minutes)\n")
    meta_p.add_run("Key Services Highlighted: ").bold = True
    meta_p.add_run("Amazon Bedrock AgentCore, AWS End User Messaging Social (WhatsApp), Amazon SES v2, AWS Pinpoint SMS v2\n")
    meta_p.add_run("Dashboard URL: ").bold = True
    meta_p.add_run("http://127.0.0.1:8000")

    doc.add_heading("Scene-by-Scene Presentation Guide", level=1)

    scenes = [
        {
            "time": "0:00 - 0:30 (30 sec)",
            "title": "Scene 1: Introduction & Problem Statement",
            "action": "Open browser at http://127.0.0.1:8000. Move mouse over the 4 top status badges (WhatsApp, SES, SMS, Bedrock).",
            "narration": (
                "Welcome to AWS NexusOps. In enterprise field operations and logistics, critical infrastructure disruptions often cause communication breakdowns. "
                "Operations leads, field engineers, and executive stakeholders are scattered across disjointed chat apps, email chains, and SMS pagers. "
                "Today, we introduce NexusOps: an autonomous omnichannel operations concierge powered by Amazon Bedrock AgentCore and AWS Communication Developer Services. "
                "NexusOps bridges WhatsApp, Amazon SES, and AWS Pinpoint SMS v2 into a unified cognitive system."
            )
        },
        {
            "time": "0:30 - 1:15 (45 sec)",
            "title": "Scene 2: Interactive WhatsApp Interaction (AWS EUM Social)",
            "action": "Click the quick button '⚡ Node 4B Failure' on the WhatsApp Mobile Phone on the left. Show the message sending and the agent's instant reply.",
            "narration": (
                "Let us observe a real-world scenario. Sarah, our Operations Director, sends an urgent message over WhatsApp via AWS End User Messaging Social. "
                "Incoming messages are processed by our backend calling boto3 social messaging. The agent analyzes the telemetry disruption on Substation Node 4B and immediately identifies the active incident ticket."
            )
        },
        {
            "time": "1:15 - 2:00 (45 sec)",
            "title": "Scene 3: Amazon Bedrock AgentCore Reasoning Traces",
            "action": "Scroll down the center column (Bedrock Agent Reasoning Traces). Point to Step 1 (Intent & get_ticket_details) and Step 2 (Autonomous Remediation & 3 tool calls).",
            "narration": (
                "Here in the center column, Amazon Bedrock AgentCore runs Anthropic Claude 3.5 Sonnet. "
                "In Step 1, Bedrock evaluates the intent and autonomously calls get_ticket_details. "
                "Upon diagnosing the disruption, the agent resolves the incident by triggering satellite failover and autonomously orchestrates all three AWS CDS channels: "
                "First, it calls send_whatsapp_message to update Sarah in real time. "
                "Second, it calls send_ses_audit_email using boto3 sesv2 to dispatch a formal executive digest. "
                "And third, it calls send_sms_urgent_alert via boto3 pinpoint-sms-voice-v2 for high-priority mobile broadcast."
            )
        },
        {
            "time": "2:00 - 2:30 (30 sec)",
            "title": "Scene 4: Multi-Channel Dispatches (SES Email & SMS Monitor)",
            "action": "Click on the right panel tabs: 1) SES Email Inbox (show styled HTML digest), 2) SMS / RCS Monitor (show yellow SMS alert card), 3) Incident Tickets (show RESOLVED status).",
            "narration": (
                "Under the SES Email tab, a styled, responsive HTML audit report has been delivered to executive stakeholders with exact incident metrics and audit timestamps. "
                "Switching to the SMS Monitor tab, the field team receives the urgent SMS alert dispatched through AWS Pinpoint SMS v2. "
                "And our ticket status is updated to Resolved in real time."
            )
        },
        {
            "time": "2:30 - 2:50 (20 sec)",
            "title": "Scene 5: Architecture & Closing",
            "action": "Show the full dashboard or display docs/architecture_diagram.svg. Provide closing statement.",
            "narration": (
                "NexusOps is built with a serverless architecture designed for planetary scale, leveraging Amazon Bedrock, AWS Lambda, DynamoDB, and AWS Communication Developer Services. "
                "By combining the conversational agility of WhatsApp, the formal accountability of Amazon SES, and the immediacy of SMS, NexusOps redefines enterprise operational resilience. "
                "Thank you!"
            )
        }
    ]

    for s in scenes:
        doc.add_heading(f"{s['title']} [{s['time']}]", level=2)
        
        p_action = doc.add_paragraph()
        p_action.add_run("🎬 Visual Action on Screen: ").bold = True
        p_action.add_run(s['action'])

        p_narr = doc.add_paragraph()
        p_narr.add_run("🎙️ Spoken Narration (What to Say): ").bold = True
        p_narr.add_run(f"\"{s['narration']}\"")
        p_narr.runs[1].font.italic = True

        doc.add_paragraph()

    doc.add_heading("Tips for a High-Scoring Video", level=2)
    tips = [
        "Keep the video under 3 minutes (Judges are not required to watch past the 3-minute mark).",
        "Keep the screen recording at 1080p Full HD resolution for sharp text.",
        "Ensure your microphone is clear and background noise is minimal.",
        "Upload to YouTube as 'Unlisted' or 'Public', and paste the URL on Devpost."
    ]
    for tip in tips:
        doc.add_paragraph(tip, style='List Bullet')

    doc.save(DOCX_PATH)
    print(f"[SUCCESS] Word document script saved to: {DOCX_PATH}")

if __name__ == "__main__":
    build_docx()
